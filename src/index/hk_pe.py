"""恒生公开月末加权市盈率快照，不推断历史分位。"""
import base64
import calendar
import logging
from datetime import date, datetime
from io import BytesIO
from zipfile import BadZipFile

import requests
from openpyxl import load_workbook

from data.cache import CacheManager
from data.collection_state import ValueResult, make_record
from data.schemas import IndexValuationData
from utils.numbers import is_positive_finite

logger = logging.getLogger(__name__)
SOURCES = {"00001.00": ("hsi", "HSI", "Hang Seng Index"),
           "02083.00": ("hstech", "HSTECH", "Hang Seng TECH Index")}


def parse_hsi_pe(content: bytes, sheet_name: str, index_name: str) -> tuple[date, float]:
    if not content.startswith(b"PK"):
        raise ValueError("恒生估值响应不是 XLSX")
    try:
        book = load_workbook(BytesIO(content), read_only=True, data_only=True)
        if sheet_name not in book.sheetnames:
            raise ValueError("恒生估值工作表不匹配")
        rows = list(book[sheet_name].values)
        if not rows or "MONTH-END WEIGHTED AVERAGE P/E RATIO" not in str(rows[0][0]):
            raise ValueError("恒生估值口径不匹配")
        header = rows[2] if len(rows) > 2 else ()
        if index_name not in header:
            raise ValueError("恒生估值指数身份不匹配")
        # HSTECH 工作表的标题在首列而数值仍在第二列，列号最低取 1。
        column = header.index(index_name) or 1
        results = []
        for row in rows[3:]:
            if len(row) <= column:
                continue
            ratio = row[column]
            if not isinstance(ratio, (int, float)) or not is_positive_finite(ratio):
                continue
            value = row[0]
            if isinstance(value, datetime):
                day = value.date()
            elif isinstance(value, date):
                day = value
            elif isinstance(value, str):
                try:
                    month = datetime.strptime(value.strip(), "%b %Y").astimezone()
                    day = date(month.year, month.month, calendar.monthrange(month.year, month.month)[1])
                except ValueError:
                    continue
            else:
                continue
            if day <= datetime.now().astimezone().date():
                results.append((day, float(ratio)))
        if not results:
            raise ValueError("恒生估值无有效月末数据")
        return max(results)
    except (KeyError, IndexError, TypeError, OSError, BadZipFile) as exc:
        raise ValueError("恒生估值文件无效") from exc


def fetch_hsi_pe(symbol: str, code: str, cache: CacheManager, *, refresh: bool = False) -> ValueResult[IndexValuationData]:
    source = SOURCES.get(code)
    if source is None:
        return ValueResult(None, make_record("恒生官方月末估值", "failed", error_kind="unavailable", error_message="未验证此指数的官方估值文件"))
    series, sheet, name = source
    url = f"https://www.hsi.com.hk/content/dam/wpb/hsil/index_operation/historical_pe/en_hk/{series}/{series}.xlsx"
    entry = cache.get_entry("hsi_pe", code, "xlsx")
    def parse(raw: bytes) -> IndexValuationData:
        day, value = parse_hsi_pe(raw, sheet, name)
        month_only = code == "02083.00"
        return IndexValuationData(symbol=symbol, date=day, pe_snapshot=value, pe_snapshot_as_of=None if month_only else day,
            pe_snapshot_period=day.strftime("%Y-%m") if month_only else "",
            pe_snapshot_basis="月末加权平均市盈率", pe_snapshot_source_url=url, valuation_valid=False,
            valuation_notes=["官方月频快照，未提供可用于长期分位的完整历史"])
    old = None
    if entry:
        try:
            old = parse(base64.b64decode(entry.value, validate=True))
        except (ValueError, TypeError):
            logger.warning("恒生估值缓存无效: %s", code)
    if old is not None and entry is not None and not refresh:
        import time
        if time.time() - entry.created_at <= entry.ttl_seconds:
            return ValueResult(old, make_record("恒生官方月末估值", "complete", cache_status="cache",
                collected_at=datetime.fromtimestamp(entry.created_at).astimezone()))
    try:
        response = requests.get(url, timeout=10)
        response.raise_for_status()
        value = parse(response.content)
        cache.put("hsi_pe", code, "xlsx", base64.b64encode(response.content).decode("ascii"), ttl_seconds=86400)
        return ValueResult(value, make_record("恒生官方月末估值", "complete", collected_at=datetime.now().astimezone()))
    except Exception as exc:  # noqa: BLE001 — 官方网络及工作簿解析边界，失败保留旧快照
        logger.warning("恒生估值采集失败 %s: %s", code, exc)
        if old is not None and entry is not None:
            return ValueResult(old, make_record("恒生官方月末估值", "partial", cache_status="stale",
                collected_at=datetime.fromtimestamp(entry.created_at).astimezone(),
                error_kind="fetch_failed", error_message="刷新失败，保留原日期估值"))
        return ValueResult(None, make_record("恒生官方月末估值", "failed", error_kind="fetch_failed", error_message="未取得有效官方估值"))
